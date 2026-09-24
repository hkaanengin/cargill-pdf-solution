"""The fixtures point at the data the rest of the suite is going to assume.

Every task from 023 onwards tests against this workbook and these ten PDFs. If
the workbook is replaced again — it has been once — or a sample goes missing,
this file fails first and says so, rather than twenty routing and lookup tests
failing for a reason none of them names.

What is asserted here is the *shape* that requirements depend on, not counts
that change with every new month's workbook.
"""

import openpyxl
import pytest

# The three filename shapes R2-R4 route by. The suite needs at least one of
# each or whole requirements go untested.
SHAPES = {
    "SGM (-> FATURA, R2)": lambda stem: stem.startswith("SGM"),
    "SUB (-> FATURA, R3)": lambda stem: stem.startswith("SUB"),
    "all digits (-> DEKONT, R4)": lambda stem: stem.isdigit(),
}


def test_every_filename_shape_has_a_sample(sample_pdfs):
    missing = [
        label for label, matches in SHAPES.items()
        if not any(matches(stem) for stem in sample_pdfs)
    ]
    assert not missing, f"No sample PDF for: {missing}"


def test_no_sample_falls_outside_the_three_shapes(sample_pdfs):
    """An unroutable name is an R5 exception, not a fixture.

    When one is wanted for testing R5 it should be made by the test, so the
    suite is never quietly relying on a stray file in samples/.
    """
    unroutable = [
        stem for stem in sample_pdfs
        if not any(matches(stem) for matches in SHAPES.values())
    ]
    assert not unroutable, f"Unroutable names in the sample directory: {unroutable}"


@pytest.mark.parametrize("position, expected_name", [(1, "FATURA"), (2, "DEKONT")])
def test_workbook_sheets_sit_where_position_says(workbook_path, position, expected_name):
    """Sheet 2 is FATURA and sheet 3 is DEKONT — R6, counting from 1.

    This asserts the *fixture*, not the app. R7 is explicit that the app never
    consults a sheet name; the suite checks it here precisely so that later
    tests can trust position without ever having to look.
    """
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        assert len(wb.worksheets) >= 3, f"Expected 3 sheets, found {wb.sheetnames}"
        assert wb.worksheets[position].title.strip() == expected_name
    finally:
        wb.close()


@pytest.mark.parametrize("position", [1, 2])
def test_both_sheets_carry_the_key_and_value_headers(workbook_path, position):
    """`Fatura No` and `PO` are found by header text, never by letter — R8."""
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        sheet = wb.worksheets[position]
        header = [
            str(cell).strip() if cell is not None else ""
            for cell in next(sheet.iter_rows(values_only=True))
        ]
        assert "Fatura No" in header, f"{sheet.title} headers: {header}"
        assert "PO" in header, f"{sheet.title} headers: {header}"
    finally:
        wb.close()
